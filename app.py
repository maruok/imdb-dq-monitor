"""
IMDB Data Quality Dashboard
Run: streamlit run app.py
"""

import io
import datetime
import streamlit as st
import plotly.graph_objects as go
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from data_loader import ensure_data
from checks import get_connection, run_all_checks, CheckResult
from investigator import investigate, continue_investigation, Investigation, meta_analyze
from prompts import load_prompt, save_prompt, DEFAULT_SYSTEM_PROMPT

st.set_page_config(
    page_title="IMDB Data Quality Monitor",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ---------------------------------------------------------------------------
# Global CSS
# ---------------------------------------------------------------------------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:opsz,wght@9..40,400;9..40,500;9..40,600;9..40,700;9..40,800&display=swap');

:root {
  --bg:        #000000;
  --bg-1:      #1c1c1e;
  --bg-2:      #2c2c2e;
  --border:    rgba(255,255,255,0.08);
  --border-2:  rgba(255,255,255,0.14);
  --text:      #ffffff;
  --text-2:    rgba(235,235,245,0.80);
  --text-3:    rgba(235,235,245,0.55);
  --teal:      #00d4aa;
  --red:       #ff6b6b;
  --blue:      #4da6ff;
  --orange:    #ff7043;
  --purple:    #a78bfa;
  --card-bg:   rgba(28,28,30,0.80);
  --card-shadow: 0 2px 12px rgba(0,0,0,0.40), 0 1px 2px rgba(0,0,0,0.50);
  --card-inset:  inset 0 1px 0 rgba(255,255,255,0.09);
  --glass-blur:  blur(18px);
  --radius:    16px;
  --radius-sm: 10px;
}

/* ── Hide Streamlit chrome ── */
#MainMenu, header, footer { visibility: hidden; }
[data-testid="stToolbar"] { display: none; }

/* ── App shell: dark ── */
html, body { background: #000 !important; font-family: 'DM Sans','Segoe UI',system-ui,sans-serif !important; }
* { font-family: 'DM Sans','Segoe UI',system-ui,sans-serif !important; }
.stApp { background: #000 !important; }
[data-testid="stAppViewContainer"] {
    background: #000 !important;
    background-image:
        radial-gradient(ellipse 120% 55% at 50% 0%,  rgba(0,212,170,0.055) 0%, transparent 65%),
        radial-gradient(ellipse 80%  40% at 80% 90%, rgba(77,166,255,0.04)  0%, transparent 60%),
        radial-gradient(ellipse 60%  30% at 15% 60%, rgba(124,77,255,0.04)  0%, transparent 55%) !important;
}
[data-testid="stMain"] { background: transparent !important; }
p, div, span, label { color: var(--text-2); }

/* ── Sidebar ── */
[data-testid="stSidebar"] {
    background: rgba(6,6,12,0.94) !important;
    border-right: 1px solid rgba(255,255,255,0.07);
}
[data-testid="stSidebar"] * { color: var(--text-2) !important; }
[data-testid="stSidebar"] label { color: var(--text-3) !important; font-size: 0.72rem !important; }
[data-testid="stSidebar"] .stSelectbox > div > div {
    background: var(--bg-1) !important;
    border-color: var(--border) !important;
    color: var(--text) !important;
}

/* ── Row spacing ── */
[data-testid="stHorizontalBlock"] {
    padding: 2px 14px; margin: 3px 0;
    align-items: center !important;
}
[data-testid="stHorizontalBlock"] > [data-testid="column"] {
    display: flex !important; flex-direction: column !important; justify-content: center !important;
}

/* ── Section containers — dark colour zones ── */
[data-testid="stVerticalBlock"]:has(.section-mark-num):not(:has(.section-mark-null)):not(:has(.section-mark-cat)) {
    background: rgba(167,139,250,0.07);
    border-radius: var(--radius); border-left: 3px solid var(--purple);
    padding: 12px 16px 16px 16px; margin-bottom: 14px;
}
[data-testid="stVerticalBlock"]:has(.section-mark-num):not(:has(.section-mark-null)):not(:has(.section-mark-cat)) [data-testid="stHorizontalBlock"] {
    background: transparent !important; border: none !important;
    border-bottom: 1px solid rgba(167,139,250,0.15) !important;
    border-radius: 0 !important; box-shadow: none !important;
}
[data-testid="stVerticalBlock"]:has(.section-mark-num):not(:has(.section-mark-null)):not(:has(.section-mark-cat)) [data-testid="stHorizontalBlock"]:last-child { border-bottom: none !important; }

[data-testid="stVerticalBlock"]:has(.section-mark-null):not(:has(.section-mark-num)):not(:has(.section-mark-cat)) {
    background: rgba(0,212,170,0.06);
    border-radius: var(--radius); border-left: 3px solid var(--teal);
    padding: 12px 16px 16px 16px; margin-bottom: 14px;
}
[data-testid="stVerticalBlock"]:has(.section-mark-null):not(:has(.section-mark-num)):not(:has(.section-mark-cat)) [data-testid="stHorizontalBlock"] {
    background: transparent !important; border: none !important;
    border-bottom: 1px solid rgba(0,212,170,0.12) !important;
    border-radius: 0 !important; box-shadow: none !important;
}
[data-testid="stVerticalBlock"]:has(.section-mark-null):not(:has(.section-mark-num)):not(:has(.section-mark-cat)) [data-testid="stHorizontalBlock"]:last-child { border-bottom: none !important; }

[data-testid="stVerticalBlock"]:has(.section-mark-cat):not(:has(.section-mark-num)):not(:has(.section-mark-null)) {
    background: rgba(77,166,255,0.06);
    border-radius: var(--radius); border-left: 3px solid var(--blue);
    padding: 12px 16px 16px 16px; margin-bottom: 14px;
}
[data-testid="stVerticalBlock"]:has(.section-mark-cat):not(:has(.section-mark-num)):not(:has(.section-mark-null)) [data-testid="stHorizontalBlock"] {
    background: transparent !important; border: none !important;
    border-bottom: 1px solid rgba(77,166,255,0.12) !important;
    border-radius: 0 !important; box-shadow: none !important;
}
[data-testid="stVerticalBlock"]:has(.section-mark-cat):not(:has(.section-mark-num)):not(:has(.section-mark-null)) [data-testid="stHorizontalBlock"]:last-child { border-bottom: none !important; }

/* ── Section header accent colours ── */
.section-header-num  { border-bottom-color: var(--purple) !important; color: var(--purple) !important; }
.section-header-null { border-bottom-color: var(--teal)   !important; color: var(--teal)   !important; }
.section-header-cat  { border-bottom-color: var(--blue)   !important; color: var(--blue)   !important; }

/* ── Summary metric cards ── */
.summary-card {
    background: var(--card-bg);
    backdrop-filter: var(--glass-blur); -webkit-backdrop-filter: var(--glass-blur);
    border-radius: var(--radius); padding: 22px 20px; text-align: center;
    box-shadow: var(--card-shadow), var(--card-inset);
    border: 1px solid rgba(255,255,255,0.09);
}
.summary-number { font-size: 2.8rem; font-weight: 800; line-height: 1; margin: 8px 0 4px 0;
                  letter-spacing: -0.04em; font-variant-numeric: tabular-nums; }
.summary-label  { font-size: 0.72rem; color: var(--text-3); text-transform: uppercase;
                  letter-spacing: 0.1em; font-weight: 700; }
.color-green { color: var(--teal); }
.color-red   { color: var(--red);  }
.color-navy  { color: var(--text); }
.color-muted { color: var(--text-3); }

/* ── Section headers ── */
.section-header {
    font-size: 0.72rem; font-weight: 800; text-transform: uppercase;
    letter-spacing: 0.14em; color: var(--text-2);
    margin: 28px 0 4px 0; padding-bottom: 8px;
    border-bottom: 2px solid rgba(255,255,255,0.10);
}

/* ── Column header labels ── */
.col-header { font-size: 0.68rem; font-weight: 700; text-transform: uppercase;
              letter-spacing: 0.1em; color: var(--text-3); }

/* ── Check row text ── */
.check-name         { font-size: 0.88rem; font-weight: 600; color: var(--text-2); }
.check-name-flagged { font-size: 0.88rem; font-weight: 700; color: var(--text);   }
.check-value-ok     { font-size: 1.05rem; font-weight: 700; color: var(--teal);   font-variant-numeric: tabular-nums; }
.check-value-flag   { font-size: 1.05rem; font-weight: 700; color: var(--red);    font-variant-numeric: tabular-nums; }
.check-range        { font-size: 0.78rem; color: var(--text-3); font-weight: 500; font-variant-numeric: tabular-nums; }

/* ── Status pills ── */
.pill { display: inline-block; padding: 4px 14px; border-radius: 20px;
        font-size: 0.7rem; font-weight: 800; letter-spacing: 0.08em; text-transform: uppercase; }
.pill-ok   { background: rgba(0,212,170,0.14);  color: #00d4aa; border: 1px solid rgba(0,212,170,0.30); }
.pill-high { background: rgba(255,107,107,0.14); color: #ff6b6b; border: 1px solid rgba(255,107,107,0.30); }
.pill-low  { background: rgba(77,166,255,0.14);  color: #4da6ff; border: 1px solid rgba(77,166,255,0.30); }

/* ── Buttons ── */
.stButton > button svg { display: none !important; }
.stButton > button {
    background: rgba(255,255,255,0.07) !important;
    color: var(--text) !important;
    border: 1px solid var(--border-2) !important;
    border-radius: 50px !important; font-size: 0.72rem !important;
    font-weight: 700 !important; padding: 5px 14px !important;
    white-space: nowrap !important; text-align: center !important;
    justify-content: center !important;
    transition: all 0.15s, transform 0.12s !important;
}
.stButton > button:hover {
    background: rgba(255,255,255,0.13) !important;
    color: var(--text) !important; border-color: rgba(255,255,255,0.25) !important;
}
.stButton > button:active { transform: scale(0.96) !important; }
[data-testid="stBaseButton-primary"], .stButton > button[kind="primary"] {
    background: var(--teal) !important; color: #000 !important;
    border: none !important; font-weight: 700 !important;
}
[data-testid="stBaseButton-primary"] *,
.stButton > button[kind="primary"] * {
    color: #000 !important;
}
[data-testid="stBaseButton-primary"]:hover, .stButton > button[kind="primary"]:hover {
    background: #00bfa0 !important;
}

/* ── Expander ── */
[data-testid="stExpander"] {
    background: var(--card-bg) !important;
    backdrop-filter: var(--glass-blur) !important; -webkit-backdrop-filter: var(--glass-blur) !important;
    border: 1px solid var(--border) !important;
    border-radius: var(--radius) !important;
    box-shadow: var(--card-shadow), var(--card-inset) !important;
}
[data-testid="stExpander"] summary {
    color: var(--text) !important; font-weight: 600 !important;
    display: flex !important; align-items: center !important; gap: 8px !important;
    overflow: hidden !important;
}
[data-testid="stExpander"] summary p { color: var(--text) !important; margin: 0 !important; }
/* hide any sr-only / visually-hidden spans that Streamlit injects before the label */
[data-testid="stExpander"] summary span:not([data-testid]) {
    position: static !important; width: auto !important; height: auto !important;
    clip: unset !important; overflow: visible !important;
    display: none !important;
}

/* ── Code blocks ── */
.stCodeBlock, pre { background: rgba(255,255,255,0.04) !important;
                    border: 1px solid var(--border) !important; border-radius: var(--radius-sm) !important; }
code { color: var(--teal) !important; }

/* ── Text area ── */
.stTextArea, .stTextArea > div, .stTextArea > label + div {
    background: var(--bg-1) !important;
}
.stTextArea textarea {
    background: var(--bg-1) !important; color: var(--text) !important;
    border: 1.5px solid var(--border-2) !important; border-radius: var(--radius-sm) !important;
    font-family: 'JetBrains Mono','Fira Code',monospace !important; font-size: 0.85rem !important;
}
.stTextArea textarea:focus { border-color: var(--teal) !important; }

/* ── Text input ── */
.stTextInput, .stTextInput > div, .stTextInput > div > div {
    background: var(--bg-1) !important;
}
.stTextInput input {
    background: var(--bg-1) !important; color: var(--text) !important;
    border: 1.5px solid var(--border-2) !important; border-radius: var(--radius-sm) !important;
}
.stTextInput input:focus { border-color: var(--teal) !important; }

/* ── Number input ── */
[data-testid="stNumberInput"], [data-testid="stNumberInput"] > div,
[data-testid="stNumberInput"] input {
    background: var(--bg-1) !important; color: var(--text) !important;
    border-color: var(--border-2) !important; border-radius: var(--radius-sm) !important;
}

/* ── Slider ── */
[data-testid="stSlider"] > div { background: transparent !important; }
[data-baseweb="slider"] { background: transparent !important; }
[data-testid="stSliderTrack"], [data-baseweb="slider"] [role="slider"] ~ div,
[data-baseweb="slider"] > div:first-child {
    background: var(--border-2) !important;
}
[data-baseweb="slider"] [role="slider"] {
    background: var(--teal) !important; border-color: var(--teal) !important;
}
[data-baseweb="slider"] [data-selected="true"] {
    background: var(--teal) !important;
}

/* ── Metrics ── */
[data-testid="metric-container"] {
    background: var(--card-bg);
    backdrop-filter: var(--glass-blur); -webkit-backdrop-filter: var(--glass-blur);
    border: 1px solid var(--border); border-radius: var(--radius-sm); padding: 10px 14px;
    box-shadow: var(--card-shadow), var(--card-inset);
}
[data-testid="stMetricValue"] { color: var(--text) !important; font-variant-numeric: tabular-nums; }
[data-testid="stMetricLabel"] { color: var(--text-3) !important; font-size: 0.72rem !important; }

/* ── Spinner / status widget ── */
[data-testid="stStatusWidget"],
[data-testid="stSpinner"],
.stSpinner {
    background: var(--card-bg) !important;
    border: 1px solid var(--border) !important;
    border-radius: var(--radius) !important;
    color: var(--text-2) !important;
    backdrop-filter: var(--glass-blur) !important;
}
[data-testid="stStatusWidget"] p,
[data-testid="stSpinner"] p,
.stSpinner p { color: var(--text-2) !important; }

/* ── Selectbox ── */
[data-testid="stSelectbox"],
[data-testid="stSelectbox"] > div,
[data-testid="stSelectbox"] > div > div,
[data-baseweb="select"],
[data-baseweb="select"] > div {
    background: var(--bg-1) !important; border-color: var(--border-2) !important;
    color: var(--text) !important; border-radius: var(--radius-sm) !important;
}
[data-baseweb="popover"], [data-baseweb="menu"] {
    background: var(--bg-2) !important; border: 1px solid var(--border-2) !important;
}
[role="option"] { background: var(--bg-2) !important; color: var(--text) !important; }
[role="option"]:hover { background: rgba(0,212,170,0.12) !important; }

/* ── Row separator ── */
.row-sep { display: none; }

/* ── Sidebar download button ── */
[data-testid="stSidebar"] .stDownloadButton > button {
    background: rgba(0,212,170,0.15) !important; color: var(--teal) !important;
    border: 1px solid rgba(0,212,170,0.30) !important; border-radius: var(--radius-sm) !important;
    font-size: 0.72rem !important; font-weight: 700 !important; width: 100% !important;
}
[data-testid="stSidebar"] .stDownloadButton > button:hover { background: rgba(0,212,170,0.25) !important; }

/* ── Alerts ── */
.stInfo    { background: rgba(77,166,255,0.10)  !important; color: #4da6ff !important;
             border-radius: var(--radius-sm) !important; border-left: 3px solid #4da6ff !important; }
.stSuccess { background: rgba(0,212,170,0.10)   !important; color: #00d4aa !important;
             border-radius: var(--radius-sm) !important; border-left: 3px solid #00d4aa !important; }
.stWarning { background: rgba(255,112,67,0.10)  !important; color: #ff7043 !important;
             border-radius: var(--radius-sm) !important; border-left: 3px solid #ff7043 !important; }
.stError   { background: rgba(255,107,107,0.10) !important; color: #ff6b6b !important;
             border-radius: var(--radius-sm) !important; border-left: 3px solid #ff6b6b !important; }
</style>
""", unsafe_allow_html=True)


if "sidebar_nav" not in st.session_state:
    st.session_state["sidebar_nav"] = "📋  Dashboard"

# ---------------------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------------------
with st.sidebar:
    st.markdown("""
    <div style='padding:8px 0 24px 0'>
        <div style='font-size:1.25rem;font-weight:800;color:#00d4aa;letter-spacing:-0.02em'>
            DQ Monitor
        </div>
        <div style='font-size:0.7rem;color:rgba(235,235,245,0.45);margin-top:2px'>IMDB Data Quality</div>
    </div>
    """, unsafe_allow_html=True)

    current_year = st.selectbox("CURRENT PERIOD", options=list(range(2024, 2009, -1)), index=0)
    n_hist = st.slider("HISTORICAL PERIODS", min_value=6, max_value=15, value=12)

    st.markdown("""
    <div style='font-size:0.72rem;color:rgba(235,235,245,0.45);line-height:1.6;margin-top:10px'>
        Flags anomalies using <b style='color:rgba(235,235,245,0.70)'>Tukey IQR fences</b>.<br>
        Current period vs preceding N years.<br>
        <b style='color:#00d4aa'>Click Investigate</b> to launch the AI agent.
    </div>
    """, unsafe_allow_html=True)



# ---------------------------------------------------------------------------
# Shared resources
# ---------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_con():
    with st.spinner("Checking data files..."):
        ensure_data()
    return get_connection()

@st.cache_data(show_spinner=False, ttl=300)
def load_checks(year: int, n: int):
    return run_all_checks(get_con(), current_year=year, n_hist=n)

if "investigations" not in st.session_state:
    st.session_state.investigations = {}
if "sql_query" not in st.session_state:
    st.session_state.sql_query = (
        "-- Available views: basics, ratings\n"
        "-- basics: tconst, titleType, primaryTitle, startYear, endYear, runtimeMinutes, genres\n"
        "-- ratings: tconst, averageRating, numVotes\n\n"
        "SELECT b.titleType, COUNT(*) AS cnt, ROUND(AVG(r.averageRating), 2) AS avg_rating\n"
        "FROM basics b JOIN ratings r USING (tconst)\n"
        "WHERE b.startYear = 2024\n"
        "GROUP BY b.titleType\n"
        "ORDER BY cnt DESC"
    )
if "aiq_prompt" not in st.session_state:
    st.session_state.aiq_prompt = load_prompt()
if "meta_suggestion" not in st.session_state:
    st.session_state.meta_suggestion = ""
if "meta_tokens" not in st.session_state:
    st.session_state.meta_tokens = (0, 0)
if "_expanded_inv" not in st.session_state:
    st.session_state._expanded_inv = ""


_TYPE_PREFIX = {"numerical": "num", "null_rate": "null", "categorical": "cat"}


def _inv_key(check: CheckResult) -> str:
    prefix = _TYPE_PREFIX.get(check.context.get("check_type", ""), "unk")
    return f"{prefix}_{check.name}_inv"


# ---------------------------------------------------------------------------
# Sparkline
# ---------------------------------------------------------------------------
def make_sparkline(check: CheckResult) -> go.Figure:
    by_year = check.context.get("by_year", {})
    curr = check.context.get("current_year")
    years = sorted(by_year.keys())
    values = [by_year[y] for y in years]
    labels = [str(y) for y in years]

    bar_colors = [
        "#ff6b6b" if (y == curr and check.flagged)
        else "#00d4aa" if y == curr
        else "#3a3a3c"
        for y in years
    ]

    fig = go.Figure()
    fig.add_trace(go.Bar(
        x=labels, y=values,
        marker_color=bar_colors,
        marker_line_width=0,
        hovertemplate="%{x}: %{y:.2f}<extra></extra>",
    ))
    fig.add_hline(y=check.fence_high, line_dash="dot", line_color="#ff7043", line_width=1.5)
    if check.fence_low > 0:
        fig.add_hline(y=check.fence_low, line_dash="dot", line_color="#4da6ff", line_width=1.5)
    fig.add_hrect(
        y0=check.fence_low, y1=check.fence_high,
        fillcolor="rgba(0,212,170,0.05)", line_width=0,
    )
    fig.update_layout(
        height=72,
        margin=dict(l=0, r=0, t=2, b=2),
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        showlegend=False,
        xaxis=dict(visible=False),
        yaxis=dict(visible=False),
        bargap=0.2,
    )
    return fig


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
def build_excel(
    checks: list,
    investigations: dict,
    current_year: int,
    n_hist: int,
    aiq_prompt: str = "",
    meta_suggestion: str = "",
) -> bytes:
    wb = openpyxl.Workbook()

    # ── Styles ──────────────────────────────────────────────────────────────
    hdr_fill   = PatternFill("solid", fgColor="1A1D35")
    flag_fill  = PatternFill("solid", fgColor="FFEBEE")
    ok_fill    = PatternFill("solid", fgColor="E8F5E9")
    grey_fill  = PatternFill("solid", fgColor="F7F8FD")
    hdr_font   = Font(bold=True, color="FFFFFF", size=10)
    flag_font  = Font(bold=True, color="C62828", size=10)
    ok_font    = Font(bold=True, color="2E7D32", size=10)
    body_font  = Font(size=10)
    title_font = Font(bold=True, size=13, color="1A1D35")
    thin       = Side(style="thin", color="D8DDF0")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)
    center     = Alignment(horizontal="center", vertical="center")
    wrap       = Alignment(wrap_text=True, vertical="top")

    def set_col_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def hdr_row(ws, row, values):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = center
            cell.border = border

    # ── Sheet 1: Summary ────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "DQ Summary"

    ws1["A1"] = "IMDB Data Quality Report"
    ws1["A1"].font = title_font
    ws1["A2"] = (
        f"Period: {current_year}  |  Baseline: {current_year - n_hist}–{current_year - 1}"
        f"  |  Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    ws1["A2"].font = Font(size=9, color="6B7094")
    ws1.row_dimensions[1].height = 22
    ws1.merge_cells("A1:G1")
    ws1.merge_cells("A2:G2")

    headers = ["Check", "Type", "Current Value", "Normal Range (Low)", "Normal Range (High)", "Status", "Direction"]
    hdr_row(ws1, 4, headers)

    type_labels = {"numerical": "Numerical", "null_rate": "Null Rate", "categorical": "Categorical"}
    for r, check in enumerate(checks, 5):
        ws1.row_dimensions[r].height = 18
        unit = f" {check.unit}" if check.unit else ""
        row_data = [
            check.name,
            type_labels.get(check.context.get("check_type", ""), ""),
            f"{check.current_val}{unit}",
            f"{check.fence_low}{unit}",
            f"{check.fence_high}{unit}",
            "FLAGGED" if check.flagged else "OK",
            check.flag_direction if check.flagged else "",
        ]
        fill = flag_fill if check.flagged else (grey_fill if r % 2 == 0 else PatternFill())
        for c, v in enumerate(row_data, 1):
            cell = ws1.cell(row=r, column=c, value=v)
            cell.font = flag_font if (check.flagged and c == 6) else (ok_font if c == 6 else body_font)
            cell.fill = fill
            cell.border = border
            cell.alignment = center if c > 1 else Alignment(vertical="center")

    set_col_widths(ws1, [38, 14, 14, 18, 18, 10, 10])

    # ── Sheet 2: Investigations ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Investigations")
    ws2["A1"] = "AI Investigation Results"
    ws2["A1"].font = title_font
    ws2.merge_cells("A1:F1")
    ws2.row_dimensions[1].height = 22

    hdr_row(ws2, 3, ["Check", "Input Tokens", "Output Tokens", "Cost (USD)", "Completed", "Summary"])

    row_num = 4
    inv_key_map = {
        k.replace("num_", "").replace("null_", "").replace("cat_", "").replace("_inv", ""): v
        for k, v in investigations.items()
    }

    for check in checks:
        # Find the matching investigation by check name
        inv_key = next(
            (k for k in investigations if check.name in k and k.endswith("_inv")), None
        )
        if not inv_key:
            continue
        inv: Investigation = investigations[inv_key]

        ws2.row_dimensions[row_num].height = 15
        summary_clean = "\n".join(
            l for l in inv.summary.split("\n") if not l.startswith("VERIFY:")
        )
        row_data = [
            check.name,
            inv.input_tokens,
            inv.output_tokens,
            f"${inv.cost_usd:.4f}",
            "Yes" if inv.completed else "No",
            summary_clean,
        ]
        for c, v in enumerate(row_data, 1):
            cell = ws2.cell(row=row_num, column=c, value=v)
            cell.font = body_font
            cell.border = border
            cell.alignment = wrap if c == 6 else Alignment(vertical="top")
            if row_num % 2 == 0:
                cell.fill = grey_fill
        ws2.row_dimensions[row_num].height = max(
            60, min(15 * (summary_clean.count("\n") + 1), 200)
        )
        row_num += 1

        # SQL steps sub-rows
        for i, step in enumerate(inv.steps, 1):
            ws2.cell(row=row_num, column=1, value=f"  Step {i} SQL").font = Font(size=9, italic=True, color="6B7094")
            cell = ws2.cell(row=row_num, column=6, value=step.sql)
            cell.font = Font(name="Courier New", size=9, color="3A3F6E")
            cell.alignment = wrap
            ws2.row_dimensions[row_num].height = 30
            row_num += 1

        # Follow-up Q&A sub-rows
        for fu_idx, fu in enumerate(inv.follow_ups, 1):
            q_cell = ws2.cell(row=row_num, column=1, value=f"  Follow-up {fu_idx} Q")
            q_cell.font = Font(size=9, italic=True, color="9399B8")
            cell = ws2.cell(row=row_num, column=6, value=fu.question)
            cell.font = Font(size=9, italic=True, color="1A1D35")
            cell.alignment = wrap
            ws2.row_dimensions[row_num].height = 25
            row_num += 1

            a_cell = ws2.cell(row=row_num, column=1, value=f"  Follow-up {fu_idx} A")
            a_cell.font = Font(size=9, italic=True, color="9399B8")
            cell = ws2.cell(row=row_num, column=6, value=fu.response)
            cell.font = Font(size=9, color="1A1D35")
            cell.alignment = wrap
            ws2.row_dimensions[row_num].height = max(
                60, min(15 * (fu.response.count("\n") + 1), 300)
            )
            row_num += 1

    set_col_widths(ws2, [38, 13, 13, 12, 10, 70])

    # ── Sheet 3: AIQ Prompt ──────────────────────────────────────────────────
    ws3 = wb.create_sheet("AIQ Prompt")
    ws3["A1"] = "AIQ Promptbook — System Prompt Used in This Report"
    ws3["A1"].font = title_font
    ws3.merge_cells("A1:B1")
    ws3.row_dimensions[1].height = 22

    meta_rows = [
        ("Report generated",   datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Current period",     str(current_year)),
        ("Historical periods", f"{current_year - n_hist}–{current_year - 1}  ({n_hist} years)"),
        ("Prompt source",      "aiq_prompt.md (file) if saved, else built-in default"),
    ]
    for r, (label, value) in enumerate(meta_rows, 3):
        ws3.cell(row=r, column=1, value=label).font  = Font(bold=True, size=10, color="6B7094")
        ws3.cell(row=r, column=2, value=value).font  = Font(size=10)
        ws3.row_dimensions[r].height = 16

    ws3.cell(row=8, column=1, value="System Prompt Text").font = Font(bold=True, size=10, color="1A1D35")
    ws3.row_dimensions[8].height = 18

    prompt_cell = ws3.cell(row=9, column=1, value=aiq_prompt or "(no prompt captured)")
    prompt_cell.font      = Font(name="Courier New", size=9, color="3A3F6E")
    prompt_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.merge_cells("A9:B9")
    ws3.row_dimensions[9].height = max(200, min(15 * aiq_prompt.count("\n"), 600))
    ws3.column_dimensions["A"].width = 60
    ws3.column_dimensions["B"].width = 40

    # ── Sheet 4: AIQ Meta-Analysis Suggestions ───────────────────────────────
    ws4 = wb.create_sheet("AIQ Meta-Analysis")
    ws4["A1"] = "AIQ Meta-Analysis — Prompt Improvement Suggestions"
    ws4["A1"].font = title_font
    ws4.merge_cells("A1:B1")
    ws4.row_dimensions[1].height = 22
    ws4["A2"] = (
        f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
        f"Based on batch run for period {current_year}"
    )
    ws4["A2"].font = Font(size=9, color="6B7094")
    ws4.merge_cells("A2:B2")

    if meta_suggestion:
        ws4.cell(row=4, column=1, value="Suggestions Text").font = Font(bold=True, size=10, color="1A1D35")
        sug_cell = ws4.cell(row=5, column=1, value=meta_suggestion)
        sug_cell.font      = Font(size=10)
        sug_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws4.merge_cells("A5:B5")
        ws4.row_dimensions[5].height = max(200, min(15 * meta_suggestion.count("\n"), 600))
    else:
        ws4.cell(row=4, column=1, value="No batch meta-analysis run yet for this session.").font = Font(size=10, color="9399B8")

    ws4.column_dimensions["A"].width = 80
    ws4.column_dimensions["B"].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Page title
# ---------------------------------------------------------------------------
st.markdown(f"""
<div style='padding:12px 0 4px 0'>
    <div style='font-size:1.65rem;font-weight:800;color:#ffffff;letter-spacing:-0.03em'>Data Quality Monitor</div>
    <div style='font-size:0.8rem;color:rgba(235,235,245,0.55);margin-top:3px'>
        Period <b style='color:#00d4aa'>{current_year}</b> &nbsp;·&nbsp;
        Baseline {current_year - n_hist}–{current_year - 1} &nbsp;·&nbsp; IMDB dataset
    </div>
</div>
""", unsafe_allow_html=True)

# Apply pending nav destination before _cur_page is read so top-nav buttons highlight immediately
if "_pending_nav" in st.session_state:
    st.session_state["sidebar_nav"] = st.session_state.pop("_pending_nav")

_cur_page = st.session_state.get("sidebar_nav", "📋  Dashboard")
_tnb1, _tnb2, _tnb3, _ = st.columns([2, 2, 2, 3])
if _tnb1.button("Dashboard", key="top_nav_db", use_container_width=True,
                type="primary" if _cur_page == "📋  Dashboard" else "secondary"):
    st.session_state["_pending_nav"] = "📋  Dashboard"
    st.rerun()
if _tnb2.button("SQL Playground", key="top_nav_sql", use_container_width=True,
                type="primary" if _cur_page == "🔍  SQL Playground" else "secondary"):
    st.session_state["_pending_nav"] = "🔍  SQL Playground"
    st.rerun()
if _tnb3.button("AIQ Promptbook", key="top_nav_aiq", use_container_width=True,
                type="primary" if _cur_page == "📝  AIQ Promptbook" else "secondary"):
    st.session_state["_pending_nav"] = "📝  AIQ Promptbook"
    st.rerun()

# Sidebar navigate + export — placed here so build_excel/load_checks are in scope
with st.sidebar:
    st.markdown("<hr style='margin:10px 0 8px 0;border-color:#2a2d4a'>", unsafe_allow_html=True)
    st.markdown("<div style='font-size:0.68rem;font-weight:700;color:#7b7fa8;letter-spacing:0.1em;text-transform:uppercase;margin-bottom:4px'>Navigate</div>", unsafe_allow_html=True)
    _nav_options = ["📋  Dashboard", "🔍  SQL Playground", "📝  AIQ Promptbook"]
    st.radio(
        "page",
        _nav_options,
        key="sidebar_nav",
        label_visibility="collapsed",
    )

    st.markdown("<div style='margin-top:10px;font-size:0.68rem;font-weight:700;color:#7b7fa8;letter-spacing:0.1em;text-transform:uppercase;margin-bottom:6px'>Export</div>", unsafe_allow_html=True)
    # Use checks already loaded by the Dashboard section (stored in session state)
    # to avoid blocking the sidebar before the loading card has a chance to render.
    _dl_checks = st.session_state.get("_checks_cache")
    if _dl_checks:
        _dl_filename = f"dq_report_{current_year}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        st.download_button(
            label="⬇  Download Report (Excel)",
            data=build_excel(_dl_checks, st.session_state.investigations, current_year, n_hist, st.session_state.aiq_prompt, st.session_state.meta_suggestion),
            file_name=_dl_filename,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    else:
        st.caption("Open the Dashboard to enable export.")

# page is read AFTER the sidebar radio has written its value to session state
page = st.session_state["sidebar_nav"]

# ===========================================================================
# DASHBOARD
# ===========================================================================
if page == "📋  Dashboard":
    _load_ph = st.empty()
    _load_ph.markdown("""
    <div style='background:rgba(28,28,30,0.80);backdrop-filter:blur(18px);
                -webkit-backdrop-filter:blur(18px);
                border:1px solid rgba(255,255,255,0.09);border-radius:16px;
                padding:48px 24px;text-align:center;margin:24px 0;
                box-shadow:0 2px 12px rgba(0,0,0,0.40);
                inset 0 1px 0 rgba(255,255,255,0.09)'>
        <div style='font-size:2rem;margin-bottom:12px'>⏳</div>
        <div style='font-weight:700;font-size:1rem;color:#ffffff'>Running statistical checks…</div>
        <div style='font-size:0.8rem;color:rgba(235,235,245,0.55);margin-top:6px'>
            Scanning IMDB data for anomalies across 3 check types
        </div>
    </div>
    """, unsafe_allow_html=True)
    checks = load_checks(current_year, n_hist)
    st.session_state["_checks_cache"] = checks
    _load_ph.empty()

    flagged   = [c for c in checks if c.flagged]
    ok_checks = [c for c in checks if not c.flagged]
    fail_rate = f"{len(flagged)/len(checks)*100:.0f}%" if checks else "0%"

    # Summary banner
    s1, s2, s3, s4 = st.columns(4)
    for col, number, label, cls in [
        (s1, len(checks),   "Checks Run",    "color-navy"),
        (s2, len(flagged),  "Flagged",       "color-red"),
        (s3, len(ok_checks),"Passed",        "color-green"),
        (s4, fail_rate,     "Failure Rate",  "color-red" if flagged else "color-green"),
    ]:
        col.markdown(f"""
        <div class="summary-card">
            <div class="summary-label">{label}</div>
            <div class="summary-number {cls}">{number}</div>
        </div>
        """, unsafe_allow_html=True)

    # ── Batch investigation run ─────────────────────────────────────────────
    if flagged:
        _bc1, _bc2, _bc3 = st.columns([2, 2, 4])
        _already_done = sum(1 for c in flagged if _inv_key(c) in st.session_state.investigations)
        _batch_label  = (
            f"Run All Flagged ({len(flagged)})"
            if _already_done == 0
            else f"Re-run All Flagged ({len(flagged)})"
        )
        _run_batch = _bc1.button(_batch_label, key="batch_run")
        if _already_done > 0:
            _bc2.markdown(
                f"<div style='font-size:0.75rem;color:#9399b8;padding-top:10px'>"
                f"{_already_done}/{len(flagged)} already investigated</div>",
                unsafe_allow_html=True,
            )

        if st.session_state.get("_batch_done_msg"):
            st.success(st.session_state.pop("_batch_done_msg"))

        if _run_batch:
            _prog = st.progress(0.0, text="Starting batch investigation…")
            for _bi, _bc in enumerate(flagged):
                _prog.progress(
                    _bi / len(flagged),
                    text=f"Investigating {_bi + 1}/{len(flagged)}: {_bc.name}",
                )
                st.session_state.investigations[_inv_key(_bc)] = investigate(
                    _bc, get_con(), st.session_state.aiq_prompt
                )
            _prog.progress(1.0, text="All investigations complete.")
            _prog.empty()
            # Rerun so sidebar download button re-renders with updated investigations
            st.session_state["_batch_done_msg"] = (
                f"Batch complete — {len(flagged)} investigations done. "
                "Go to **AIQ Promptbook** to review results and request prompt suggestions."
            )
            st.rerun()

    def send_to_playground(q: str):
        st.session_state.sql_query = q

    def render_check_row(check: CheckResult, section_key: str):
        key     = f"{section_key}_{check.name}"
        inv_key = key + "_inv"

        c_name, c_chart, c_val, c_range, c_status, c_action = st.columns(
            [2.4, 2.2, 0.9, 1.6, 0.85, 1.05]
        )

        name_cls = "check-name-flagged" if check.flagged else "check-name"
        c_name.markdown(f"<div class='{name_cls}'>{check.name}</div>", unsafe_allow_html=True)

        c_chart.plotly_chart(
            make_sparkline(check), use_container_width=True,
            config={"displayModeBar": False},
        )

        unit = f" {check.unit}" if check.unit else ""
        val_cls = "check-value-flag" if check.flagged else "check-value-ok"
        c_val.markdown(
            f"<div class='{val_cls}'>{check.current_val}{unit}</div>",
            unsafe_allow_html=True,
        )
        c_range.markdown(
            f"<div class='check-range'>"
            f"{check.fence_low}{unit} – {check.fence_high}{unit}</div>",
            unsafe_allow_html=True,
        )

        if check.flagged:
            pill_cls = "pill-high" if check.flag_direction == "HIGH" else "pill-low"
            c_status.markdown(
                f"<span class='pill {pill_cls}'>{check.flag_direction}</span>",
                unsafe_allow_html=True,
            )
            if c_action.button("Investigate", key=key):
                with st.spinner("AI investigating..."):
                    inv = investigate(check, get_con())
                    st.session_state.investigations[inv_key] = inv
                st.session_state._expanded_inv = inv_key
        else:
            c_status.markdown(
                "<span class='pill pill-ok'>OK</span>",
                unsafe_allow_html=True,
            )

        # Investigation result
        if inv_key in st.session_state.investigations:
            inv = st.session_state.investigations[inv_key]
            _is_expanded = (inv_key == st.session_state.get("_expanded_inv", ""))
            with st.expander(f"Investigation: {check.name}", expanded=_is_expanded):
                t1, t2, t3, t4 = st.columns(4)
                t1.metric("Input tokens",  f"{inv.input_tokens:,}")
                t2.metric("Output tokens", f"{inv.output_tokens:,}")
                t3.metric("Total tokens",  f"{inv.total_tokens:,}")
                t4.metric("Cost",          f"${inv.cost_usd:.4f}")
                st.markdown("---")

                if inv.steps:
                    st.markdown("#### Reasoning trace")
                    for i, step in enumerate(inv.steps, 1):
                        st.markdown(f"**Step {i}**")
                        if step.reasoning:
                            st.info(step.reasoning)
                        sc1, sc2 = st.columns([5, 1])
                        sc1.code(step.sql, language="sql")
                        if sc2.button("▶ Run", key=f"{inv_key}_step_{i}"):
                            send_to_playground(step.sql)
                            st.info("Sent to SQL Playground — click the tab above.")
                        st.code(step.result)
                    st.markdown("---")

                st.markdown("#### Summary")
                lines        = inv.summary.split("\n")
                summary_text = "\n".join(
                    l for l in lines
                    if not l.startswith("VERIFY:")
                    and not l.strip().startswith("ACTION REQUIRED:")
                    and not l.strip().startswith("NO ACTION NEEDED:")
                )
                _verdict = next(
                    (l.strip() for l in lines
                     if l.strip().startswith("ACTION REQUIRED:") or l.strip().startswith("NO ACTION NEEDED:")),
                    None,
                )

                if _verdict:
                    if _verdict.startswith("ACTION REQUIRED:"):
                        st.error(f"🔴 {_verdict}")
                    else:
                        st.success(f"✅ {_verdict}")

                if inv.completed:
                    st.success(summary_text)
                else:
                    st.warning(summary_text)

                # ── Follow-up chat ──────────────────────────────────────────
                st.markdown("---")
                st.markdown("#### Follow-up Questions")

                for fu_idx, fu in enumerate(inv.follow_ups):
                    with st.chat_message("user"):
                        st.markdown(fu.question)
                    with st.chat_message("assistant"):
                        if fu.steps:
                            for si, step in enumerate(fu.steps, 1):
                                with st.expander(f"Step {si}", expanded=False):
                                    if step.reasoning:
                                        st.info(step.reasoning)
                                    fsc1, fsc2 = st.columns([5, 1])
                                    fsc1.code(step.sql, language="sql")
                                    if fsc2.button("▶ Run", key=f"{inv_key}_fu{fu_idx}_s{si}"):
                                        send_to_playground(step.sql)
                                        st.info("Sent to SQL Playground — click the tab above.")
                                    st.code(step.result)
                        st.markdown(fu.response)
                        st.caption(
                            f"Tokens: {fu.input_tokens + fu.output_tokens:,}"
                            f"  ·  Cost: ${fu.cost_usd:.4f}"
                        )

                with st.form(key=f"{inv_key}_fu_form", clear_on_submit=True):
                    fu_col1, fu_col2 = st.columns([6, 1])
                    fu_question = fu_col1.text_input(
                        "follow-up",
                        placeholder="Ask a follow-up question, e.g. Which specific titles drove the change?",
                        label_visibility="collapsed",
                    )
                    fu_submit = fu_col2.form_submit_button("Send →", type="primary")

                if fu_submit and fu_question.strip():
                    with st.spinner("AI investigating follow-up..."):
                        continue_investigation(inv, fu_question.strip(), get_con())
                    st.session_state._expanded_inv = inv_key
                    st.rerun()

        st.markdown("<hr class='row-sep'>", unsafe_allow_html=True)

    def render_section(title: str, section_checks: list[CheckResult], key_prefix: str, section_class: str):
        if not section_checks:
            return
        with st.container():
            st.markdown(f"<div class='section-mark {section_class}'></div>", unsafe_allow_html=True)
            st.markdown(
                f"<div class='section-header section-header-{key_prefix}'>{title}</div>",
                unsafe_allow_html=True,
            )
            h = st.columns([2.4, 2.2, 0.9, 1.6, 0.85, 1.05])
            for col, label in zip(h, ["Check", "12-period trend", "Current", "Normal range", "Status", "Action"]):
                col.markdown(f"<div class='col-header'>{label}</div>", unsafe_allow_html=True)
            for check in sorted(section_checks, key=lambda c: (not c.flagged, c.name)):
                render_check_row(check, key_prefix)

    numerical  = [c for c in checks if c.context.get("check_type") == "numerical"]
    nulls      = [c for c in checks if c.context.get("check_type") == "null_rate"]
    categorical= [c for c in checks if c.context.get("check_type") == "categorical"]

    render_section("Numerical Variables",      numerical,   "num",  "section-mark-num")
    render_section("Null Rate Monitoring",     nulls,       "null", "section-mark-null")
    render_section("Categorical Distribution", categorical, "cat",  "section-mark-cat")



# ===========================================================================
# SQL PLAYGROUND
# ===========================================================================
elif page == "🔍  SQL Playground":
    st.markdown("<div class='section-header'>SQL Playground</div>", unsafe_allow_html=True)
    st.markdown(
        "<div style='font-size:0.82rem;color:#9399b8;margin-bottom:16px'>"
        "Query the IMDB data directly. Available views: "
        "<code style='background:rgba(0,212,170,0.15);padding:2px 8px;border-radius:4px;color:#00d4aa;border:1px solid rgba(0,212,170,0.30)'>basics</code> &nbsp;"
        "<code style='background:rgba(0,212,170,0.15);padding:2px 8px;border-radius:4px;color:#00d4aa;border:1px solid rgba(0,212,170,0.30)'>ratings</code>"
        "</div>",
        unsafe_allow_html=True,
    )

    examples = {
        "Title counts by type and year": (
            "SELECT b.startYear, b.titleType, COUNT(*) AS cnt\n"
            "FROM basics b JOIN ratings r USING (tconst)\n"
            "WHERE b.startYear BETWEEN 2015 AND 2024\n"
            "GROUP BY b.startYear, b.titleType\n"
            "ORDER BY b.startYear, cnt DESC"
        ),
        "Top 20 highest rated movies (1000+ votes)": (
            "SELECT b.primaryTitle, b.startYear, r.averageRating, r.numVotes\n"
            "FROM basics b JOIN ratings r USING (tconst)\n"
            "WHERE b.titleType = 'movie' AND r.numVotes >= 1000\n"
            "ORDER BY r.averageRating DESC\n"
            "LIMIT 20"
        ),
        "Null rate for runtimeMinutes by year": (
            "SELECT b.startYear, COUNT(*) AS total,\n"
            "       SUM(CASE WHEN b.runtimeMinutes IS NULL THEN 1 ELSE 0 END) AS nulls,\n"
            "       ROUND(100.0 * SUM(CASE WHEN b.runtimeMinutes IS NULL THEN 1 ELSE 0 END) / COUNT(*), 1) AS null_pct\n"
            "FROM basics b JOIN ratings r USING (tconst)\n"
            "WHERE b.titleType = 'movie' AND b.startYear BETWEEN 2012 AND 2024\n"
            "GROUP BY b.startYear ORDER BY b.startYear"
        ),
        "Genre distribution for a year": (
            "SELECT b.genres, COUNT(*) AS cnt,\n"
            "       ROUND(100.0 * COUNT(*) / SUM(COUNT(*)) OVER (), 2) AS pct\n"
            "FROM basics b JOIN ratings r USING (tconst)\n"
            "WHERE b.startYear = 2024 AND b.genres IS NOT NULL\n"
            "GROUP BY b.genres ORDER BY cnt DESC LIMIT 20"
        ),
    }

    with st.expander("Example queries"):
        for label, sql in examples.items():
            ec1, ec2 = st.columns([5, 1])
            ec1.markdown(
                f"<div style='font-size:0.85rem;font-weight:600;color:var(--text-2,rgba(235,235,245,0.80));padding:6px 0'>{label}</div>",
                unsafe_allow_html=True,
            )
            if ec2.button("Use", key=f"ex_{label}"):
                st.session_state.sql_query = sql
                st.rerun()
            st.code(sql, language="sql")

    query = st.text_area("Write your SQL query:", value=st.session_state.sql_query, height=200)

    rc1, rc2, _ = st.columns([1.2, 1, 6])
    run_clicked = rc1.button("▶  Run query", type="primary")
    if rc2.button("Clear"):
        st.session_state.sql_query = ""
        st.rerun()

    if run_clicked and query.strip():
        try:
            cursor = get_con().execute(query)
            rows   = cursor.fetchall()
            cols_n = [desc[0] for desc in cursor.description]
            if not rows:
                st.info("Query returned no rows.")
            else:
                data = {cols_n[i]: [r[i] for r in rows] for i in range(len(cols_n))}
                st.markdown(
                    f"<div style='font-size:0.78rem;color:#4caf50;margin:8px 0'>"
                    f"✓ {len(rows)} row(s) returned</div>",
                    unsafe_allow_html=True,
                )
                st.dataframe(data, use_container_width=True)
        except Exception as e:
            st.error(f"SQL Error: {e}")


# ===========================================================================
# AIQ PROMPTBOOK
# ===========================================================================
elif page == "📝  AIQ Promptbook":
    st.markdown("<div class='section-header'>AIQ Promptbook</div>", unsafe_allow_html=True)
    st.markdown(
        "<div style='font-size:0.82rem;color:#9399b8;margin-bottom:20px'>"
        "Edit the system prompt given to the AI investigation agent. "
        "Changes take effect immediately on the next investigation. "
        "The prompt is saved to <code style='background:rgba(0,212,170,0.15);padding:2px 8px;"
        "border-radius:4px;color:#00d4aa;border:1px solid rgba(0,212,170,0.30)'>aiq_prompt.md</code> in the project folder."
        "</div>",
        unsafe_allow_html=True,
    )

    # ── Prompt editor ────────────────────────────────────────────────────────
    _pa, _pb, _pc, _ = st.columns([2, 2, 2, 3])
    _save_clicked  = _pa.button("Save Changes", key="aiq_save", use_container_width=True)
    _reset_clicked = _pb.button("Reset to Default", key="aiq_reset", use_container_width=True)
    if _pc.button("Reload from File", key="aiq_reload", use_container_width=True):
        st.session_state.aiq_prompt = load_prompt()
        st.rerun()

    if _reset_clicked:
        st.session_state.aiq_prompt = DEFAULT_SYSTEM_PROMPT
        st.rerun()

    edited_prompt = st.text_area(
        "System prompt:",
        value=st.session_state.aiq_prompt,
        height=420,
        label_visibility="collapsed",
    )

    if _save_clicked:
        st.session_state.aiq_prompt = edited_prompt
        save_prompt(edited_prompt)
        st.success("Saved to aiq_prompt.md — all future investigations will use this prompt.")

    # ── Prompt improvement suggestions ───────────────────────────────────────
    st.markdown("---")
    st.markdown(
        "<div class='section-header'>Prompt Improvement Suggestions</div>",
        unsafe_allow_html=True,
    )

    _n_inv = len(st.session_state.investigations)
    _ra, _rb, _ = st.columns([3, 3, 3])
    _ra.markdown(
        (
            f"<div style='font-size:0.82rem;color:#6b7094;padding-top:10px'>"
            f"{_n_inv} investigation(s) available for review</div>"
        ) if _n_inv > 0 else (
            "<div style='font-size:0.82rem;color:#9399b8;padding-top:10px'>"
            "Run at least one investigation from the Dashboard first.</div>"
        ),
        unsafe_allow_html=True,
    )
    _run_meta = _rb.button(
        "Review & suggest improvements",
        key="aiq_meta_run",
        use_container_width=True,
        disabled=_n_inv == 0,
    )
    if _run_meta:
            _summaries = []
            for inv in st.session_state.investigations.values():
                if not inv.summary:
                    continue
                text = inv.summary
                for fu in inv.follow_ups:
                    if fu.response:
                        text += (
                            f"\n\nFollow-up question from analyst: {fu.question}"
                            f"\nFollow-up analysis: {fu.response}"
                        )
                _summaries.append(text)
            with st.spinner(f"Reviewing {len(_summaries)} investigation(s)…"):
                _suggestion, _in_tok, _out_tok = meta_analyze(
                    _summaries, edited_prompt
                )
            st.session_state.meta_suggestion = _suggestion
            st.session_state.meta_tokens     = (_in_tok, _out_tok)
            st.rerun()

    if st.session_state.meta_suggestion:
        st.markdown("<div style='margin-top:8px'></div>", unsafe_allow_html=True)
        _in_tok, _out_tok = st.session_state.meta_tokens
        _cost = (_in_tok * 3.0 + _out_tok * 15.0) / 1_000_000
        st.caption(f"Tokens: {_in_tok + _out_tok:,}  ·  Cost: ${_cost:.4f}  ·  Included in Excel export")

        # Show only the analysis sections (Common Patterns, Gaps, Suggestions) — not the full revised prompt
        _full_meta = st.session_state.meta_suggestion
        _revised_marker = "## Revised Full Prompt"
        _display_meta = (
            _full_meta[:_full_meta.find(_revised_marker)].strip()
            if _revised_marker in _full_meta
            else _full_meta
        )
        st.markdown(_display_meta)

        _btn_col1, _btn_col2, _ = st.columns([3, 3, 3])
        if _btn_col1.button("Append suggestions to prompt", key="aiq_append"):
            # Extract only the Suggested Prompt Additions section — not the analyst commentary
            _full = st.session_state.meta_suggestion
            _start = "## Suggested Prompt Additions"
            _end   = "## Revised Full Prompt"
            _additions_only = (
                _full[_full.find(_start) + len(_start) : _full.find(_end)].strip()
                if _start in _full and _end in _full
                else (_full[_full.find(_start) + len(_start):].strip() if _start in _full else _full)
            )
            _ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
            appended = (
                edited_prompt.rstrip()
                + f"\n\n# --- AIQ suggested additions ({_ts}) ---\n"
                + _additions_only
            )
            st.session_state.aiq_prompt = appended
            save_prompt(appended)
            st.rerun()

        # Revised full prompt section
        if _revised_marker in _full_meta:
            _revised_prompt = _full_meta[_full_meta.find(_revised_marker) + len(_revised_marker):].strip()
            with st.expander("Replace full prompt with AI-revised version", expanded=False):
                st.caption(
                    "The AI has rewritten the full prompt incorporating all improvements. "
                    "Review before applying — edits are saved immediately."
                )
                _edited_revision = st.text_area(
                    "Revised prompt:",
                    value=_revised_prompt,
                    height=340,
                    key="aiq_revised_prompt",
                    label_visibility="collapsed",
                )
                if st.button("Replace full prompt with this revision", key="aiq_replace_full"):
                    _ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
                    _dated = _edited_revision.rstrip() + f"\n\n# Revised by meta-analysis: {_ts}"
                    st.session_state.aiq_prompt = _dated
                    save_prompt(_dated)
                    st.success(f"Full prompt replaced ({_ts}).")
                    st.rerun()
